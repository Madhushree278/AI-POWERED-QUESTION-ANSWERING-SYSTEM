from pathlib import Path
import pymupdf
import docx

def create_sample_pdf(filepath: Path):
    doc = pymupdf.open()
    
    # Page 1: Title & Chapter 1
    page1 = doc.new_page()
    text1 = """Artificial Intelligence and Machine Learning Handbook
Author: Dr. Alan Turing Institute

Chapter 1: Introduction to Artificial Intelligence

Artificial Intelligence (AI) is defined as the simulation of human intelligence in machines that are programmed to think and learn. 
The core objective of AI is to develop systems capable of performing tasks that typically require human cognition, such as visual perception, speech recognition, decision-making, and language translation.

The Turing Test, introduced by Alan Turing in 1950, remains a foundational benchmark for assessing whether a machine can exhibit intelligent behavior indistinguishable from that of a human. Early AI systems relied heavily on symbolic logic, expert rule-based engines, and heuristic search algorithms. Modern AI has shifted decisively toward data-driven statistical models, probabilistic reasoning, and deep artificial neural networks.
"""
    page1.insert_text((50, 60), text1, fontsize=12)

    # Page 2: Chapter 2
    page2 = doc.new_page()
    text2 = """Chapter 2: Machine Learning Fundamentals

Machine learning (ML) is a subset of artificial intelligence focused on building applications that learn from data and improve their accuracy over time without being explicitly programmed.

The three primary paradigms of machine learning are:
1. Supervised Learning: Algorithms are trained on labeled datasets where each input is paired with the correct output. Common algorithms include linear regression, logistic regression, support vector machines (SVM), decision trees, and random forests. Key applications include spam detection, image classification, and medical diagnosis.
2. Unsupervised Learning: Algorithms discover hidden patterns, groupings, or features within unlabeled datasets. Key techniques include k-means clustering, hierarchical clustering, and principal component analysis (PCA).
3. Reinforcement Learning: An agent learns by interacting with an environment through trial and error, receiving rewards for desirable actions and penalties for undesirable ones. Algorithms include Q-learning and Deep Q-Networks (DQN).
"""
    page2.insert_text((50, 60), text2, fontsize=12)

    # Page 3: Chapter 3
    page3 = doc.new_page()
    text3 = """Chapter 3: Deep Neural Networks and Transformers

Deep Learning is a specialized subfield of machine learning inspired by the structure and function of the human brain. Deep learning utilizes artificial neural networks with multiple hidden layers (hence "deep") to automatically extract high-level representations from raw input data.

A Perceptron is the fundamental building unit of an artificial neural network, calculating a weighted sum of inputs and applying an activation function such as ReLU or Sigmoid.
Backpropagation is the primary algorithm used to train neural networks, calculating the gradient of the loss function with respect to each weight via the chain rule of calculus and updating parameters using stochastic gradient descent (SGD).

The Transformer architecture, introduced by Vaswani et al. in 2017 with the seminal paper "Attention Is All You Need", replaced recurrent mechanisms with multi-head self-attention. Transformers enable massive parallelization across large corpora, giving rise to modern Large Language Models (LLMs) such as GPT, Gemini, and Claude.
"""
    page3.insert_text((50, 60), text3, fontsize=12)

    # Page 4: Chapter 4
    page4 = doc.new_page()
    text4 = """Chapter 4: Ethics, Alignment, and AI Safety

As artificial intelligence systems become increasingly capable and pervasive across critical societal domains, addressing safety, fairness, transparency, and alignment becomes paramount.

Algorithmic bias occurs when machine learning models inherit, reflect, or amplify systemic prejudices present in their training data. For example, facial recognition algorithms trained primarily on lighter-skinned faces exhibit significantly higher error rates for individuals with darker skin tones.

AI Alignment refers to the challenge of ensuring that artificial intelligence systems reliably pursue goals aligned with human values, ethics, and intent. Failure of alignment can lead to reward hacking or catastrophic unintended consequences. Interpretability techniques, such as SHAP (SHapley Additive exPlanations) and attention map visualization, help researchers audit model decisions.
"""
    page4.insert_text((50, 60), text4, fontsize=12)

    # Add TOC (Table of Contents)
    doc.set_toc([
        [1, "Chapter 1: Introduction to Artificial Intelligence", 1],
        [1, "Chapter 2: Machine Learning Fundamentals", 2],
        [1, "Chapter 3: Deep Neural Networks and Transformers", 3],
        [1, "Chapter 4: Ethics, Alignment, and AI Safety", 4]
    ])

    doc.save(str(filepath))
    doc.close()
    print(f"Created PDF: {filepath}")

def create_sample_docx(filepath: Path):
    doc = docx.Document()
    doc.add_heading("Cloud Computing and AWS Essentials", level=0)
    
    doc.add_heading("Chapter 1: Introduction to Cloud Infrastructure", level=1)
    doc.add_paragraph(
        "Cloud computing is the on-demand delivery of IT resources over the Internet with pay-as-you-go pricing. "
        "Instead of buying, owning, and maintaining physical data centers and servers, organizations can access technology services, "
        "such as computing power, storage, and databases, on an as-needed basis from a cloud provider like Amazon Web Services (AWS)."
    )
    doc.add_paragraph(
        "The three primary cloud computing models are:\n"
        "1. Infrastructure as a Service (IaaS): Provides basic building blocks for cloud IT, typically providing access to networking features, computers (virtual or on dedicated hardware), and data storage space.\n"
        "2. Platform as a Service (PaaS): Removes the need for organizations to manage underlying infrastructure (hardware and operating systems) and allows focus on deployment and management of applications.\n"
        "3. Software as a Service (SaaS): Provides a complete product that is run and managed by the service provider, such as web-based email."
    )

    doc.add_heading("Chapter 2: Core AWS Services", level=1)
    doc.add_paragraph(
        "Amazon Web Services (AWS) provides over 200 fully featured services from data centers globally.\n"
        "Core computing services include:\n"
        "- Amazon EC2 (Elastic Compute Cloud): Provides secure, resizable compute capacity in the cloud.\n"
        "- AWS Lambda: Serverless compute engine that runs code in response to events without provisioning servers.\n"
        "- Amazon S3 (Simple Storage Service): Scalable object storage designed for 99.999999999% (11 9s) of data durability.\n"
        "- Amazon RDS (Relational Database Service): Makes it easy to set up, operate, and scale relational databases such as PostgreSQL, MySQL, and Aurora."
    )

    doc.save(str(filepath))
    print(f"Created DOCX: {filepath}")

if __name__ == "__main__":
    samples_dir = Path("sample_documents")
    samples_dir.mkdir(exist_ok=True)
    create_sample_pdf(samples_dir / "AI_and_ML_Handbook.pdf")
    create_sample_docx(samples_dir / "AWS_Cloud_Essentials.docx")
